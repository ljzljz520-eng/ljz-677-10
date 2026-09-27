"""财政补贴名单导入平台

经办人上传补贴名单（姓名/身份证号/银行卡号/金额/乡镇）→
系统格式校验 + 重复校验 → 分批上送财政平台 →
平台退回异常生成错误清单（可导出、可重试）。
乡镇账号数据隔离，只能查看本乡镇的导入任务。
"""
import csv
import io
import os
import random
from datetime import datetime
from functools import wraps

from flask import (Flask, Response, abort, flash, redirect, render_template,
                   request, session, url_for)
from werkzeug.security import check_password_hash

from db import close_db, get_db
from fiscal_client import MockFiscalPlatform
from validators import (compute_id_check_code, luhn_check_digit,
                        validate_amount, validate_bank_card, validate_id_card,
                        validate_name)

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'dev-secret-key-change-me')
app.config['MAX_CONTENT_LENGTH'] = 8 * 1024 * 1024   # 上传文件最大 8MB
app.teardown_appcontext(close_db)

BATCH_SIZE = 50     # 每批上送笔数
PAGE_SIZE = 20      # 明细分页大小
MAX_ROWS = 5000     # 单文件最大数据行数

STATUS_LABELS = {'VALID': '校验通过', 'INVALID': '校验失败',
                 'SUCCESS': '上送成功', 'FAILED': '上送失败'}
TASK_STATUS_LABELS = {'VALIDATED': '待提交', 'COMPLETED': '已完成'}
BATCH_STATUS_LABELS = {'SUBMITTED': '已上送', 'ACCEPTED': '全部受理',
                       'PARTIAL': '部分受理', 'REJECTED': '整批退回'}

# 上传文件列名别名（表头自动识别）
COLUMN_ALIASES = {
    'name':      ['姓名', '补贴对象', '名称', '户主姓名', '收款人'],
    'id_card':   ['身份证号', '身份证', '证件号码', '公民身份号码', '身份证号码'],
    'bank_card': ['银行卡号', '银行卡', '卡号', '银行账号', '账号'],
    'amount':    ['金额', '发放金额', '补贴金额', '补贴标准'],
    'township':  ['乡镇', '乡镇名称', '所属乡镇', '乡镇街道'],
}


# ---------------- 基础工具 ----------------

@app.context_processor
def inject_globals():
    return dict(user=current_user(),
                STATUS_LABELS=STATUS_LABELS,
                TASK_STATUS_LABELS=TASK_STATUS_LABELS,
                BATCH_STATUS_LABELS=BATCH_STATUS_LABELS,
                BATCH_SIZE=BATCH_SIZE)


def current_user():
    if 'user_id' not in session:
        return None
    return get_db().execute('SELECT * FROM users WHERE id=?',
                            (session['user_id'],)).fetchone()


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        return view(*args, **kwargs)
    return wrapped


def get_task_or_403(task_id, user):
    """按数据权限取任务：乡镇账号只能访问本乡镇任务"""
    task = get_db().execute('SELECT * FROM import_tasks WHERE id=?',
                            (task_id,)).fetchone()
    if task is None:
        abort(404)
    if user['role'] != 'admin' and task['township'] != user['township']:
        abort(403)
    return task


# ---------------- 登录 ----------------

@app.route('/login', methods=['GET', 'POST'])
def login():
    if 'user_id' in session:
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        user = get_db().execute('SELECT * FROM users WHERE username=?',
                                (username,)).fetchone()
        if user and check_password_hash(user['password_hash'], password):
            session.clear()
            session['user_id'] = user['id']
            return redirect(url_for('dashboard'))
        flash('用户名或密码错误', 'error')
    return render_template('login.html')


@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))


# ---------------- 任务列表（按乡镇隔离） ----------------

@app.route('/')
@login_required
def dashboard():
    user = current_user()
    db = get_db()
    if user['role'] == 'admin':
        tasks = db.execute(
            'SELECT t.*, u.display_name FROM import_tasks t'
            ' JOIN users u ON u.id = t.user_id ORDER BY t.id DESC LIMIT 100').fetchall()
        stats = db.execute(
            "SELECT (SELECT COUNT(*) FROM import_tasks) AS task_count,"
            " (SELECT COUNT(*) FROM subsidy_records) AS record_count,"
            " (SELECT COUNT(*) FROM subsidy_records WHERE status='SUCCESS') AS success_count,"
            " (SELECT COUNT(*) FROM subsidy_records WHERE status IN ('INVALID','FAILED')) AS error_count,"
            " (SELECT COALESCE(SUM(amount),0) FROM subsidy_records WHERE status='SUCCESS') AS success_amount"
        ).fetchone()
    else:
        tasks = db.execute(
            'SELECT t.*, u.display_name FROM import_tasks t'
            ' JOIN users u ON u.id = t.user_id WHERE t.township=?'
            ' ORDER BY t.id DESC LIMIT 100', (user['township'],)).fetchall()
        stats = db.execute(
            "SELECT (SELECT COUNT(*) FROM import_tasks WHERE township=?) AS task_count,"
            " (SELECT COUNT(*) FROM subsidy_records WHERE township=?) AS record_count,"
            " (SELECT COUNT(*) FROM subsidy_records WHERE township=? AND status='SUCCESS') AS success_count,"
            " (SELECT COUNT(*) FROM subsidy_records WHERE township=? AND status IN ('INVALID','FAILED')) AS error_count,"
            " (SELECT COALESCE(SUM(amount),0) FROM subsidy_records WHERE township=? AND status='SUCCESS') AS success_amount",
            (user['township'],) * 5).fetchone()
    return render_template('dashboard.html', tasks=tasks, stats=stats)


# ---------------- 文件解析 ----------------

def cell_str(v):
    """单元格转字符串（Excel 数字单元格去掉 .0 尾巴）"""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_upload(file_storage):
    filename = file_storage.filename or ''
    ext = os.path.splitext(filename)[1].lower()
    if ext == '.csv':
        return parse_csv(file_storage)
    if ext == '.xlsx':
        return parse_xlsx(file_storage)
    raise ValueError('仅支持 .csv 或 .xlsx 文件')


def parse_csv(fs):
    raw = fs.read()
    text = None
    for enc in ('utf-8-sig', 'gbk'):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError('文件编码无法识别，请使用 UTF-8 或 GBK 编码')
    try:
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=',\t;')
    except csv.Error:
        dialect = csv.excel
    rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    rows = [r for r in rows if any(r)]
    return rows_to_dicts(rows)


def parse_xlsx(fs):
    from openpyxl import load_workbook
    wb = load_workbook(fs, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        cells = [cell_str(v) for v in row]
        if any(cells):
            rows.append(cells)
    return rows_to_dicts(rows)


def rows_to_dicts(rows):
    """首行表头识别列位置，其余行转 dict 并做基础清洗"""
    if not rows:
        raise ValueError('文件内容为空')
    header = rows[0]
    col_map = {}
    for field, aliases in COLUMN_ALIASES.items():
        for idx, h in enumerate(header):
            if h in aliases:
                col_map[field] = idx
                break
    names = {'name': '姓名', 'id_card': '身份证号', 'bank_card': '银行卡号',
             'amount': '金额', 'township': '乡镇'}
    missing = [names[f] for f in names if f not in col_map]
    if missing:
        raise ValueError('缺少必需列：' + '、'.join(missing))

    result = []
    for r in rows[1:]:
        def get(field):
            idx = col_map[field]
            return r[idx] if idx < len(r) else ''
        result.append({
            'name': get('name').strip(),
            'id_card': get('id_card').strip().upper(),
            'bank_card': get('bank_card').strip().replace(' ', ''),
            'amount': get('amount').strip(),
            'township': get('township').strip(),
        })
    if not result:
        raise ValueError('文件中没有数据行')
    return result


# ---------------- 校验与建任务 ----------------

def validate_row(row, row_no, user, seen, db):
    """单条记录校验：格式校验 + 重复校验。返回 (错误列表, 金额Decimal或None)"""
    errors = []

    ok, msg = validate_name(row['name'])
    if not ok:
        errors.append(msg)

    id_ok, msg = validate_id_card(row['id_card'])
    if not id_ok:
        errors.append(msg)

    ok, msg = validate_bank_card(row['bank_card'])
    if not ok:
        errors.append(msg)

    ok, msg, amount_value = validate_amount(row['amount'])
    if not ok:
        errors.append(msg)

    township = row['township']
    if not township:
        errors.append('乡镇不能为空')
    elif user['role'] != 'admin' and township != user['township']:
        errors.append(f'乡镇"{township}"与当前账号所属乡镇（{user["township"]}）不符')

    # 重复校验（仅对格式合法的身份证）
    if id_ok:
        id_card = row['id_card']
        if id_card in seen:                       # 文件内重复
            errors.append(f'与文件内第 {seen[id_card]} 行身份证号重复')
        else:
            seen[id_card] = row_no
        dup = db.execute(                         # 库内重复（待上送/已发放）
            "SELECT t.task_no FROM subsidy_records r"
            " JOIN import_tasks t ON t.id = r.task_id"
            " WHERE r.id_card=? AND r.status IN ('VALID','SUCCESS') LIMIT 1",
            (id_card,)).fetchone()
        if dup:
            errors.append(f'身份证号已存在于任务 {dup["task_no"]}，疑似重复发放')
    return errors, amount_value


def create_task(user, filename, rows):
    """校验全部行并落库，返回 (任务ID, 通过数, 失败数)"""
    db = get_db()
    task_no = 'RW%s%02d%02d' % (datetime.now().strftime('%Y%m%d%H%M%S'),
                                user['id'], random.randint(0, 99))
    cur = db.execute(
        'INSERT INTO import_tasks (task_no, user_id, township, filename)'
        ' VALUES (?,?,?,?)',
        (task_no, user['id'], user['township'], filename))
    task_id = cur.lastrowid

    seen = {}
    valid_count = invalid_count = 0
    for idx, row in enumerate(rows, start=1):
        errors, amount_value = validate_row(row, idx, user, seen, db)
        if errors:
            invalid_count += 1
            status, err_msg = 'INVALID', '；'.join(errors)
        else:
            valid_count += 1
            status, err_msg = 'VALID', None
        db.execute(
            'INSERT INTO subsidy_records'
            ' (task_id, row_number, name, id_card, bank_card, amount, township, status, error_message)'
            ' VALUES (?,?,?,?,?,?,?,?,?)',
            (task_id, idx, row['name'], row['id_card'], row['bank_card'],
             float(amount_value) if amount_value is not None else 0.0,
             row['township'] or user['township'], status, err_msg))
    db.execute('UPDATE import_tasks SET total_count=?, valid_count=?, invalid_count=?'
               ' WHERE id=?', (len(rows), valid_count, invalid_count, task_id))
    db.commit()
    return task_id, valid_count, invalid_count


@app.route('/upload', methods=['GET', 'POST'])
@login_required
def upload():
    user = current_user()
    if request.method == 'POST':
        f = request.files.get('file')
        if not f or not f.filename:
            flash('请选择要上传的文件', 'error')
            return redirect(url_for('upload'))
        try:
            rows = parse_upload(f)
        except ValueError as exc:
            flash(f'文件解析失败：{exc}', 'error')
            return redirect(url_for('upload'))
        if len(rows) > MAX_ROWS:
            flash(f'单个文件最多 {MAX_ROWS} 行数据', 'error')
            return redirect(url_for('upload'))
        task_id, valid_count, invalid_count = create_task(user, f.filename, rows)
        flash(f'导入完成：共 {len(rows)} 行，校验通过 {valid_count} 行，校验失败 {invalid_count} 行',
              'success' if invalid_count == 0 else 'error')
        return redirect(url_for('task_detail', task_id=task_id))
    return render_template('upload.html')


@app.route('/template.csv')
@login_required
def download_template():
    """下载导入模板（含一行合法示例数据）"""
    user = current_user()
    id17 = '11010119900307123'
    id_card = id17 + compute_id_check_code(id17)
    bank_card = '622202020011223' + luhn_check_digit('622202020011223')
    township = '城关镇' if user['role'] == 'admin' else user['township']
    out = io.StringIO()
    out.write('﻿')  # BOM，保证 Excel 打开不乱码
    w = csv.writer(out)
    w.writerow(['姓名', '身份证号', '银行卡号', '发放金额', '乡镇'])
    w.writerow(['张三', id_card, bank_card, '1500.00', township])
    return Response(out.getvalue(), mimetype='text/csv',
                    headers={'Content-Disposition': 'attachment; filename=import_template.csv'})


# ---------------- 分批上送 ----------------

def do_submit(task_id):
    """将任务中状态为 VALID 的记录分批上送财政平台，返回 (成功数, 失败数)"""
    db = get_db()
    records = db.execute(
        "SELECT * FROM subsidy_records WHERE task_id=? AND status='VALID'"
        ' ORDER BY row_number', (task_id,)).fetchall()
    if not records:
        return 0, 0

    platform = MockFiscalPlatform()
    base_no = db.execute('SELECT COALESCE(MAX(batch_no),0) m FROM batches WHERE task_id=?',
                         (task_id,)).fetchone()['m']
    now = lambda: datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    total_ok = total_fail = 0

    for i in range(0, len(records), BATCH_SIZE):
        chunk = records[i:i + BATCH_SIZE]
        batch_no = base_no + i // BATCH_SIZE + 1
        cur = db.execute(
            'INSERT INTO batches (task_id, batch_no, record_count, submitted_at)'
            ' VALUES (?,?,?,?)', (task_id, batch_no, len(chunk), now()))
        batch_id = cur.lastrowid

        payload = [dict(row_number=r['row_number'], name=r['name'],
                        id_card=r['id_card'], bank_card=r['bank_card'],
                        amount=r['amount']) for r in chunk]
        results = platform.submit_batch(batch_no, payload)

        n_fail = 0
        for r, res in zip(chunk, results):
            if res['success']:
                db.execute("UPDATE subsidy_records SET status='SUCCESS',"
                           ' batch_id=?, error_message=NULL WHERE id=?',
                           (batch_id, r['id']))
            else:
                db.execute("UPDATE subsidy_records SET status='FAILED',"
                           ' batch_id=?, error_message=? WHERE id=?',
                           (batch_id, res['message'], r['id']))
                n_fail += 1
        total_ok += len(chunk) - n_fail
        total_fail += n_fail

        bstatus = ('ACCEPTED' if n_fail == 0
                   else 'REJECTED' if n_fail == len(chunk) else 'PARTIAL')
        db.execute('UPDATE batches SET status=?, finished_at=?, message=? WHERE id=?',
                   (bstatus, now(), f'成功 {len(chunk) - n_fail} 笔，失败 {n_fail} 笔',
                    batch_id))

    db.execute("UPDATE import_tasks SET status='COMPLETED',"
               ' success_count=success_count+?, failed_count=? WHERE id=?',
               (total_ok, total_fail, task_id))
    db.commit()
    return total_ok, total_fail


@app.route('/tasks/<int:task_id>/submit', methods=['POST'])
@login_required
def submit_task(task_id):
    user = current_user()
    get_task_or_403(task_id, user)
    ok, fail = do_submit(task_id)
    if ok == 0 and fail == 0:
        flash('没有可上送的记录（需状态为"校验通过"）', 'error')
    else:
        flash(f'上送完成：成功 {ok} 笔，失败 {fail} 笔',
              'success' if fail == 0 else 'error')
    return redirect(url_for('task_detail', task_id=task_id))


@app.route('/tasks/<int:task_id>/retry', methods=['POST'])
@login_required
def retry_task(task_id):
    """失败记录重新置为待提交并再次上送"""
    user = current_user()
    get_task_or_403(task_id, user)
    db = get_db()
    n = db.execute("UPDATE subsidy_records SET status='VALID', error_message=NULL,"
                   " batch_id=NULL WHERE task_id=? AND status='FAILED'",
                   (task_id,)).rowcount
    db.commit()
    if n == 0:
        flash('没有失败记录可重试', 'error')
        return redirect(url_for('task_detail', task_id=task_id))
    ok, fail = do_submit(task_id)
    flash(f'重试完成：成功 {ok} 笔，仍失败 {fail} 笔',
          'success' if fail == 0 else 'error')
    return redirect(url_for('task_detail', task_id=task_id))


# ---------------- 任务详情 / 错误清单 ----------------

@app.route('/tasks/<int:task_id>')
@login_required
def task_detail(task_id):
    user = current_user()
    task = get_task_or_403(task_id, user)
    db = get_db()
    status = request.args.get('status', '')
    page = max(request.args.get('page', 1, type=int) or 1, 1)

    where, params = 'task_id=?', [task_id]
    if status in STATUS_LABELS:
        where += ' AND status=?'
        params.append(status)
    total = db.execute(f'SELECT COUNT(*) c FROM subsidy_records WHERE {where}',
                       params).fetchone()['c']
    records = db.execute(
        f'SELECT * FROM subsidy_records WHERE {where}'
        ' ORDER BY row_number LIMIT ? OFFSET ?',
        params + [PAGE_SIZE, (page - 1) * PAGE_SIZE]).fetchall()
    batches = db.execute('SELECT * FROM batches WHERE task_id=? ORDER BY batch_no',
                         (task_id,)).fetchall()
    pages = max((total + PAGE_SIZE - 1) // PAGE_SIZE, 1)
    return render_template('task_detail.html', task=task, records=records,
                           batches=batches, status=status, page=page,
                           pages=pages, total=total)


def get_error_records(task_id):
    return get_db().execute(
        "SELECT * FROM subsidy_records WHERE task_id=?"
        " AND status IN ('INVALID','FAILED') ORDER BY row_number",
        (task_id,)).fetchall()


@app.route('/tasks/<int:task_id>/errors')
@login_required
def error_list(task_id):
    user = current_user()
    task = get_task_or_403(task_id, user)
    return render_template('errors.html', task=task,
                           records=get_error_records(task_id))


@app.route('/tasks/<int:task_id>/errors/export')
@login_required
def export_errors(task_id):
    """导出错误清单 CSV（带 BOM，Excel 可直接打开）"""
    user = current_user()
    get_task_or_403(task_id, user)
    out = io.StringIO()
    out.write('﻿')
    w = csv.writer(out)
    w.writerow(['行号', '姓名', '身份证号', '银行卡号', '金额', '乡镇',
                '错误类型', '错误原因'])
    for r in get_error_records(task_id):
        w.writerow([r['row_number'], r['name'], r['id_card'], r['bank_card'],
                    f"{r['amount']:.2f}", r['township'],
                    '格式校验' if r['status'] == 'INVALID' else '平台退回',
                    r['error_message']])
    return Response(
        out.getvalue(), mimetype='text/csv',
        headers={'Content-Disposition':
                 f'attachment; filename=error_list_task_{task_id}.csv'})


# ---------------- 错误页 ----------------

@app.errorhandler(403)
def forbidden(e):
    return render_template('error.html', code=403,
                           message='无权访问：该任务属于其他乡镇'), 403


@app.errorhandler(404)
def not_found(e):
    return render_template('error.html', code=404, message='页面不存在'), 404


@app.errorhandler(413)
def too_large(e):
    flash('文件超过大小限制（8MB）', 'error')
    return redirect(url_for('upload'))


if __name__ == '__main__':
    app.run(host='127.0.0.1', port=5000, debug=False)
